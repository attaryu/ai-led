"""Endpoint export semua section approved jadi 1 .docx."""
import io
from typing import Optional

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from docx import Document

from lib.db import conn
from lib.outline import get_sections

router = APIRouter(prefix="/export", tags=["export"])


@router.get("")
def export_docx(model: str = Query(..., description="Model whose drafts to export"),
                only_approved: bool = Query(True)):
    """Assemble drafts sesuai urutan outline -> .docx download."""
    outline = get_sections()
    section_ids = [s["id"] for s in outline]
    titles = {s["id"]: s["title"] for s in outline}

    with conn() as con:
        if only_approved:
            rows = con.execute(
                "SELECT section_id, draft_text FROM drafts "
                "WHERE model=? AND status='approved'", (model,)).fetchall()
        else:
            rows = con.execute(
                "SELECT section_id, draft_text FROM drafts WHERE model=?",
                (model,)).fetchall()
    draft_map = {r["section_id"]: r["draft_text"] for r in rows}

    if not draft_map:
        raise HTTPException(404, "Belum ada draft untuk diekspor")

    doc = Document()
    doc.add_heading("Laporan Evaluasi Diri Program Studi", level=0)
    doc.add_paragraph(f"Generated via: {model}")

    for sid in section_ids:
        if sid not in draft_map:
            continue
        doc.add_heading(titles[sid], level=1)
        for para in draft_map[sid].split("\n\n"):
            para = para.strip()
            if para:
                doc.add_paragraph(para)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    fname = f"LED_{model.replace(':','_').replace('/','_')}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
