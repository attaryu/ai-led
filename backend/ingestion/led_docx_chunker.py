"""
LED (Laporan Evaluasi Diri) .docx chunker.

Extracts text from a LED .docx paragraph-by-paragraph, tracking paragraph
offsets and heading-style positions, then applies the same sliding-window
chunking strategy used for PDF (ingestion.led_chunker._sliding_window_chunks).

Unlike the PDF path, section headings are detected from the docx's own
paragraph style ("Heading 1", "Heading 2", ...) rather than regex — this
works even when the heading text itself has no BAB/C.x marker prefix (e.g.
LED docx exported by the AI-LED writing system drops the marker, keeping
only the bare title).
"""

from __future__ import annotations

import uuid
from typing import Optional

from docx import Document

from ingestion.led_chunker import _PARA_BREAK, _sliding_window_chunks

# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def chunk_led_docx(docx_path: str):
    """Extract text from a LED .docx and chunk it with overlap.

    Steps:
    1. Walk paragraphs in document order via python-docx.
    2. Concatenate paragraph text with paragraph-break markers, recording
       char offsets per paragraph and per heading-styled paragraph.
    3. Sliding-window chunk the combined text (shared with the PDF path).
    4. Resolve paragraph-index positions and headings for each chunk window.
    """
    combined_text, para_offsets, heading_offsets = _extract_docx(docx_path)
    sorted_heading_offsets = sorted(heading_offsets.keys())

    def heading_resolver(_text: str, char_start: int, _char_end: int) -> Optional[str]:
        return _heading_at_or_before(char_start, sorted_heading_offsets, heading_offsets)

    return _sliding_window_chunks(
        combined_text,
        position_offsets=para_offsets,
        heading_resolver=heading_resolver,
        position_field="para",
    )


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------


def _extract_docx(docx_path: str) -> tuple[str, dict[int, int], dict[int, str]]:
    """Return (combined_text, {char_offset: paragraph_index}, {char_offset: heading_text}).

    para_offsets records the char offset where each paragraph starts, keyed
    to its 0-based paragraph index. heading_offsets records only paragraphs
    whose style is "Heading N" (the document title, added via add_heading at
    level 0 -> style "Title", is intentionally excluded — it is not a LED
    section).
    """
    doc = Document(docx_path)
    parts: list[str] = []
    para_offsets: dict[int, int] = {}
    heading_offsets: dict[int, str] = {}
    current_offset = 0

    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    for para_idx, para in enumerate(paragraphs):
        para_offsets[current_offset] = para_idx

        style_name = (para.style.name or "") if para.style else ""
        if style_name.startswith("Heading"):
            heading_offsets[current_offset] = para.text.strip()

        raw = para.text
        parts.append(raw)
        current_offset += len(raw)

        if para_idx < len(paragraphs) - 1:
            parts.append(_PARA_BREAK)
            current_offset += len(_PARA_BREAK)

    return "".join(parts), para_offsets, heading_offsets


def _heading_at_or_before(
    offset: int,
    sorted_heading_offsets: list[int],
    heading_offsets: dict[int, str],
) -> Optional[str]:
    """Return the heading text whose paragraph most closely precedes `offset`.

    Binary search, mirrors ingestion.led_chunker._value_at_offset but returns
    None (not 0) when no heading precedes the offset at all.
    """
    if not sorted_heading_offsets:
        return None
    lo, hi = 0, len(sorted_heading_offsets) - 1
    result: Optional[str] = None
    while lo <= hi:
        mid = (lo + hi) // 2
        if sorted_heading_offsets[mid] <= offset:
            result = heading_offsets[sorted_heading_offsets[mid]]
            lo = mid + 1
        else:
            hi = mid - 1
    return result
