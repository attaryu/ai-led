"""Tests for excel reader and writer modules."""

import os
import shutil
from pathlib import Path

import openpyxl
import pytest

from excel.reader import IndicatorRow, read_indicator_rows, _COL_SCORE, _COL_JUST
from excel.writer import write_assessment_result

LAPORAN_PATH = os.environ.get(
    "LAPORAN_TEST_FILE",
    "/home/arielsulton/Documents/Stargazing Project/VScode Project/tugas_akhir/AMI/docs/reference/Laporan_AK.xlsx",
)


@pytest.fixture(scope="session")
def indicators():
    if not Path(LAPORAN_PATH).exists():
        pytest.skip(f"Laporan_AK not available at: {LAPORAN_PATH}")
    return read_indicator_rows(LAPORAN_PATH)


# ── 1 ──────────────────────────────────────────────────────────────────────────
def test_read_returns_list_when_file_missing(tmp_path):
    """read_indicator_rows must return [] for a missing file, not raise."""
    result = read_indicator_rows(str(tmp_path / "does_not_exist.xlsx"))
    assert isinstance(result, list)
    assert result == []


# ── 2 ──────────────────────────────────────────────────────────────────────────
def test_read_indicator_count(indicators):
    assert len(indicators) == 78, (
        f"Expected exactly 78 indicators, got {len(indicators)}. "
        "Update this assertion if the workbook legitimately changed."
    )


# ── 3 ──────────────────────────────────────────────────────────────────────────
def test_indicator_row_fields(indicators):
    for ind in indicators:
        assert isinstance(ind, IndicatorRow)
        assert isinstance(ind.row_index, int) and ind.row_index >= 0
        assert isinstance(ind.indicator_no, int) and ind.indicator_no >= 1
        assert isinstance(ind.kriteria, str) and len(ind.kriteria) > 0


# ── 4 ──────────────────────────────────────────────────────────────────────────
def test_indicator_nos_unique(indicators):
    nos = [ind.indicator_no for ind in indicators]
    assert len(nos) == len(set(nos)), "Duplicate indicator_no values found"


# ── 5 ──────────────────────────────────────────────────────────────────────────
def test_write_score_persists(indicators, tmp_path):
    """Write score=3 to first indicator's judge row, re-read, verify col E equals 3."""
    copy_path = str(tmp_path / "laporan_copy.xlsx")
    shutil.copy(LAPORAN_PATH, copy_path)

    first = indicators[0]
    write_assessment_result(copy_path, first.row_index, score=3, justification="")

    wb = openpyxl.load_workbook(copy_path, data_only=True)
    from excel.reader import _find_kertas_kerja_sheet
    ws = _find_kertas_kerja_sheet(wb)
    assert ws is not None
    openpyxl_row = first.row_index + 1   # judge row, 1-based
    written_score = ws.cell(row=openpyxl_row, column=_COL_SCORE + 1).value
    wb.close()

    assert written_score == 3, f"Expected score 3 in judge row col E, got {written_score}"


# ── 6 ──────────────────────────────────────────────────────────────────────────
def test_write_justification_persists(indicators, tmp_path):
    """Write justification text to judge row col H, re-read, verify."""
    copy_path = str(tmp_path / "laporan_copy2.xlsx")
    shutil.copy(LAPORAN_PATH, copy_path)

    first = indicators[0]
    expected_text = "Bukti menunjukkan bahwa program studi telah memenuhi standar."
    write_assessment_result(copy_path, first.row_index, score=4, justification=expected_text)

    wb = openpyxl.load_workbook(copy_path, data_only=True)
    from excel.reader import _find_kertas_kerja_sheet
    ws = _find_kertas_kerja_sheet(wb)
    assert ws is not None
    openpyxl_row = first.row_index + 1
    written_text = ws.cell(row=openpyxl_row, column=_COL_JUST + 1).value
    wb.close()

    assert written_text == expected_text, f"Expected justification text, got: {written_text!r}"


# ── 7 ──────────────────────────────────────────────────────────────────────────
def test_row_indices_are_monotonically_increasing(indicators):
    """Indicators must appear in sheet order."""
    indices = [ind.row_index for ind in indicators]
    assert indices == sorted(indices), "Indicator row_index values are not in sheet order"
