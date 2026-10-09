"""1 file LKPS (xlsx/xls) -> list chunk dict siap diembed + di-index ke
vectorstore/pgvector_lkps_store.py. Reuse penuh evidence_excel.py
(extract_tables/drop_first_col_only_rows/df_to_konten) -- algoritma yang sama
sudah dipakai apps/ai untuk parsing Excel Bukti Dukung (evidence_pipeline.py).

SATU SHEET = SATU CHUNK (beda dari evidence_pipeline.py::_process_excel yang
pecah per sub-tabel) -- port dari Dimas's process_lkps_file (lib/extraction/
pipeline.py): kalau 1 sheet berisi >1 sub-tabel (+ tabel "Info"), semua
digabung jadi satu blok teks per sheet, dipisah baris kosong antar sub-tabel.
Keputusan produk yang sama dipertahankan: retrieval (search_lkps) mengembalikan
konteks 1 sheet utuh per hit, bukan potongan sub-tabel lepas -- LLM generator
lebih mudah menyusun narasi dari 1 sheet lengkap drpd menyambung beberapa
chunk terpisah.

Semua sheet diindeks, termasuk navigasi ("Menu"/"Daftar Tabel") -- sama
dengan dimas; _LKPS_RAG_DISTANCE_MAX (0.38) yang menyaring sheet navigasi.
"""
from __future__ import annotations

import re
from pathlib import Path

from .evidence_chunking import clean_entity_evidence
from .evidence_excel import df_to_konten, drop_first_col_only_rows, extract_tables

# Awalan "Tabel 4.f)" / "Tabel 2.A.1" di judul tabel -- nomor tabel tidak membantu
# pencarian makna, hanya menambah noise ke embedding.
_RE_AWALAN_TABEL = re.compile(r"^\s*Tabel\s+[\w.\-]+\)?\s*", re.IGNORECASE)
# "Bagian-3" / "Bagian 3" -> "Bagian": nomor bagian sama noise-nya dengan nomor tabel.
_RE_NOMOR_BAGIAN = re.compile(r"\b(Bagian)[\s\-]*\d+\b", re.IGNORECASE)
_KOLOM_DIABAIKAN = {"no", "no.", "nomor"}


def bersihkan_judul_tabel(label: str) -> str:
    """'Tabel 4.f) Luaran ... DTPS | Tabel 4.f) Bagian-3 Teknologi Tepat Guna, Produk'
    -> 'Luaran ... DTPS Bagian Teknologi Tepat Guna Produk'. Bagian yang sama
    persis setelah dibersihkan hanya ditulis sekali."""
    bagian: list[str] = []
    for potong in str(label).split("|"):
        potong = _RE_AWALAN_TABEL.sub("", potong.strip())
        potong = _RE_NOMOR_BAGIAN.sub(r"\1", potong)
        potong = re.sub(r"[,;:]", "", potong)
        potong = re.sub(r"\s+", " ", potong).strip()
        if potong and potong not in bagian:
            bagian.append(potong)
    return " ".join(bagian)


def nama_kolom_bersih(tables) -> list[str]:
    """Nama kolom semua sub-tabel, berurutan, tanpa duplikat; kolom nomor urut
    dan kolom tanpa nama (Unnamed) dibuang. Header bertingkat ('Induk | Anak')
    ditulis 'Induk Anak'."""
    out: list[str] = []
    for df in tables:
        for kolom in df.columns:
            nama = " ".join(dict.fromkeys(
                re.sub(r"\s+", " ", str(b)).strip() for b in str(kolom).split("|")
            ))
            if not nama or nama.lower() in _KOLOM_DIABAIKAN or nama.lower().startswith("unnamed"):
                continue
            if nama not in out:
                out.append(nama)
    return out


def teks_embedding(label: str, tables) -> str:
    """Teks yang di-EMBED untuk satu sheet: hanya judul tabel + nama kolom.
    Isi nilai tabel TIDAK ikut -- itu tetap di `text` (yang dikirim ke LLM)."""
    judul = bersihkan_judul_tabel(label)
    kolom = nama_kolom_bersih(tables)
    teks = f"{judul}."
    if kolom:
        teks += f"\nberisi: {', '.join(kolom)}."
    return teks


def process_lkps_file(xlsx_path: Path) -> list[dict]:
    """Return [{chunk_id, sheet, text}, ...] -- satu dict per sheet yang
    punya isi. `chunk_id` = nama sheet (unik per dokumen, sesuai
    @@unique([documentId, chunkId]) di LkpsChunkEmbedding)."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, data_only=True)
    out: list[dict] = []
    for sheet in wb.sheetnames:
        try:
            result = extract_tables(wb, sheet_name=sheet)
        except Exception:
            continue

        tables = drop_first_col_only_rows(result["tables"])
        table_name = result.get("name")
        label = table_name or sheet

        parts: list[str] = []
        for t_idx, df in enumerate(tables):
            if df.empty:
                continue
            konten = df_to_konten(df)
            if not konten.strip():
                continue
            heading = f"{label} Sub-tabel {t_idx + 1}" if len(tables) > 1 else label
            parts.append(f"{heading}\n{konten}")

        df_info = result.get("info")
        if df_info is not None and not df_info.empty:
            konten_info = df_to_konten(df_info)
            if konten_info.strip():
                parts.append(f"Info\n{konten_info}")

        if not parts:
            continue

        out.append({
            "chunk_id": sheet,
            "sheet": sheet,
            "text": "\n\n".join(parts),
            # Hanya untuk embedding (judul tabel + nama kolom, bersih); `text` tetap
            # judul + seluruh isi tabel dan itulah yang dikirim ke LLM.
            "embed_text": teks_embedding(label, tables),
            # port dari Dimas's entity label -- dipakai kalau chunk ini perlu
            # ditampilkan sbg sitasi/preview, sama konvensi evidence_pipeline.py.
            "entity": clean_entity_evidence(f"Sheet {sheet} {label}"),
        })
    return out
