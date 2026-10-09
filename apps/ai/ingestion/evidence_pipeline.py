"""Facade: one BuktiDukung file -> list of chunk dicts, ready to embed.

Adapted from a sibling prototype's lib/extraction/pipeline.py. doc/docx/
ppt/pptx are converted to PDF first (LibreOffice headless, needs
libreoffice-writer/-impress in the image). jpg/jpeg/png intentionally not
ported (no pipeline anywhere extracts text from a standalone image —
MinerU is a PDF/office tool, not a generic OCR tool): a real gap, not a
bug — see worker.py's ingestion handler for how callers are expected to
handle UnsupportedFileError (GAGAL, not silently SIAP-with-zero-chunks, so
the UI doesn't claim a file is AI-searchable when it never got chunked).
"""

from __future__ import annotations

import re
from pathlib import Path

from .evidence_chunking import chunk_evidence_markdown, clean_entity_evidence  # noqa: F401 -- chunk_evidence_markdown dipertahankan utk rujukan
from .evidence_chunking_v2 import (
    MAX_TOKENS,
    MERGE_CHUNKS,
    STRUCTURAL_HEADINGS,
    chunk_document,
)
from .evidence_excel import df_to_konten, drop_first_col_only_rows, extract_tables
from . import mineru_service
from .evidence_mineru import convert_to_pdf, mineru_extract, mineru_extract_paged
from .watermark import remove_watermark_text

CONVERTIBLE_EXTS = {".doc", ".docx", ".ppt", ".pptx"}
EXCEL_EXTS = {".xlsx", ".xls"}
TEXT_EXTS = {".txt"}
PDF_EXTS = {".pdf"}
CSV_EXTS = {".csv"}


class UnsupportedFileError(Exception):
    """Extension we don't (yet) know how to extract text from."""


def _file_kind(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in CONVERTIBLE_EXTS:
        return "convert"
    if ext in EXCEL_EXTS:
        return "excel"
    if ext in CSV_EXTS:
        return "csv"
    if ext in TEXT_EXTS:
        return "text"
    if ext in PDF_EXTS:
        return "pdf"
    try:
        with open(path, "rb") as f:
            head = f.read(5)
        if head == b"%PDF-":
            return "pdf"
    except OSError:
        pass
    return "unknown"


# Worker menyimpan berkas sebagai "<uuid>-<nama asli>.pdf" (lihat worker.py: local_path =
# work_dir / Path(file_url).name), jadi `file_path.stem` berawalan UUID.
_UUID_AWAL = re.compile(r"^[0-9a-fA-F]{8}-(?:[0-9a-fA-F]{4}-){3}[0-9a-fA-F]{12}-")


def _nama_dokumen(filename: str) -> str:
    """Nama yang layak dibaca model: tanpa awalan UUID, garis bawah jadi spasi.

    ponytail: chunking v2 menaruh `Dokumen: {nama}` di SETIAP chunk (tercatat di
    embed_text). Tanpa pembersihan, UUID ~13 token ikut ter-embed dan memakan
    anggaran MAX_TOKENS pada tiap chunk. chunk_id tetap memakai `filename` asli
    supaya keunikannya (UNIQUE evidence_id, chunk_id) tidak berubah."""
    bersih = _UUID_AWAL.sub("", filename).replace("_", " ").strip()
    return bersih or filename


def _bersihkan_nul(teks: str) -> str:
    """Buang karakter NUL (0x00) dari teks hasil ekstraksi.

    PostgreSQL menolak string yang memuat NUL ("A string literal cannot contain NUL
    (0x00) characters"), sehingga SATU chunk bernul menggagalkan INSERT dan seluruh
    berkas berstatus GAGAL -- padahal chunk lainnya baik. Kasus nyata: "6 Renstra PENS
    2020-2024.pdf" menghasilkan 1.160 NUL berderet tepat setelah teks "Formula:" (rumus
    yang gagal didekode MinerU), membuat 6 dari 72 chunk tidak bisa disimpan. Dibuang di
    sini, sebelum chunking, supaya deretan NUL juga tidak memakan jatah token chunk."""
    return teks.replace("\x00", "")


def _process_pdf(pdf_path: Path, filename: str, mineru_out_dir: Path) -> list[dict]:
    """MinerU -> markdown -> chunking v2 (evidence_chunking_v2.chunk_document).

    ponytail: pindah dari chunk_evidence_markdown (evidence_chunking.py) ke v2.
    Yang lama tidak punya batas ukuran chunk -- besar chunk mengikuti panjang isi
    di bawah satu heading, sehingga ada chunk 51.176 karakter, padahal
    agent/tools/evidence_tool.py memotong `[:1000]` sebelum mengirim ke LLM
    (67% isi bukti dukung tidak pernah terbaca). v2 membatasi MAX_TOKENS=1000
    bersatuan token, menggabung chunk bertetangga selama total <= MAX_TOKENS, dan memecah tabel per baris
    dengan header diulang.

    KOLOM KELUARAN TIDAK BERUBAH: {entity, text, chunk_id, confidence}, jadi
    skema database dan index_evidence_chunks tidak ikut tersentuh.
      entity     <- `context` (judul heading induk) + nama berkas, pola yang sama
                    dengan versi lama lewat clean_entity_evidence.
      text       <- `embed_text` (sudah memuat awalan "Dokumen: <nama>" dan
                    konteks heading), yaitu teks yang memang dimaksudkan untuk
                    di-embed.
      confidence <- "LOW" bila chunk ditandai oversize/sangat_pendek oleh v2,
                    "HIGH" selainnya -- menjaga arti kolom lama (kualitas chunk),
                    bukan sekadar diisi None.
    """
    # ponytail: mineru_extract_paged -- sisipkan penanda `<!-- page N of M -->` yang
    # dibutuhkan chunking v2 (buang header/footer berulang, page_start/page_end).
    # md_path = mineru_extract(pdf_path, mineru_out_dir)
    if mineru_service.service_enabled():
        # MinerU 4.x di container `mineru` (GPU lokal ATAU mineru.net, ditentukan MINERU_MODE
        # di service). Keluarannya sudah memuat penanda halaman -- sama dengan notebook.
        # raw_md = remove_watermark_text(mineru_service.parse_pdf(pdf_path))
        raw_md = remove_watermark_text(_bersihkan_nul(mineru_service.parse_pdf(pdf_path)))
    else:
        # Cadangan: MinerU 3.4.5 in-process + penanda halaman buatan sendiri.
        md_path = mineru_extract_paged(pdf_path, mineru_out_dir)
        # raw_md = remove_watermark_text(md_path.read_text(encoding="utf-8"))
        raw_md = remove_watermark_text(_bersihkan_nul(md_path.read_text(encoding="utf-8")))

    # ev_chunks = chunk_evidence_markdown(raw_md, filename)
    nama = _nama_dokumen(filename)
    v2_chunks, _log, _meta = chunk_document(
        # raw_md, nama, MAX_TOKENS, MIN_TOKENS, STRUCTURAL_HEADINGS   # notebook lama: gabung hanya < MIN_TOKENS
        # MERGE_CHUNKS: penggabungan baru dari ekstraksi_mineru (1).ipynb -- gabung ke chunk
        # berikutnya selama total token <= MAX_TOKENS (3 chunk atau lebih boleh).
        raw_md, nama, MAX_TOKENS, MERGE_CHUNKS, STRUCTURAL_HEADINGS
    )

    hasil = []
    for i, c in enumerate(v2_chunks):
        konteks = c.get("context") or ""
        ragu = {"oversize", "sangat_pendek"} & set(c.get("flags") or [])
        hasil.append(
            {
                "entity": clean_entity_evidence(f"{nama} {konteks}".strip()),
                "text": c["embed_text"],
                "chunk_id": f"evidence__{filename}__{i:04d}",
                "confidence": "LOW" if ragu else "HIGH",
            }
        )
    return hasil


def _process_excel(xlsx_path: Path, filename: str) -> list[dict]:
    from openpyxl import load_workbook

    # Loaded once and reused across sheets — extract_tables() used to take a
    # filepath and reload+reparse the whole workbook per sheet, which made
    # multi-sheet LKPS uploads (tens of sheets) crawl.
    wb = load_workbook(xlsx_path, data_only=True)
    out: list[dict] = []
    for sheet in wb.sheetnames:
        try:
            result = extract_tables(wb, sheet_name=sheet)
        except Exception:
            continue
        result["tables"] = drop_first_col_only_rows(result["tables"])
        table_name = result.get("name")
        for t_idx, df in enumerate(result["tables"]):
            if df.empty:
                continue
            konten = df_to_konten(df)
            if not konten.strip():
                continue
            label = table_name or sheet
            entity = clean_entity_evidence(f"{filename.replace('_', ' ')} {sheet} {label} Sub-tabel {t_idx + 1}")
            out.append({
                "entity": entity,
                "text": konten,
                "chunk_id": f"evidence__{filename}__{sheet}__{t_idx:02d}",
                "confidence": None,
            })
        df_info = result.get("info")
        if df_info is not None and not df_info.empty:
            konten_info = df_to_konten(df_info)
            if konten_info.strip():
                entity = clean_entity_evidence(f"{filename.replace('_', ' ')} {sheet} Info")
                out.append({
                    "entity": entity,
                    "text": konten_info,
                    "chunk_id": f"evidence__{filename}__{sheet}__info",
                    "confidence": None,
                })
    return out


def _process_csv(csv_path: Path, filename: str) -> list[dict]:
    import pandas as pd

    df = pd.read_csv(csv_path)
    if df.empty:
        return []
    konten = df_to_konten(df)
    if not konten.strip():
        return []
    entity = clean_entity_evidence(filename.replace("_", " "))
    return [{
        "entity": entity,
        "text": konten,
        "chunk_id": f"evidence__{filename}__0000",
        "confidence": None,
    }]


def _process_text(txt_path: Path, filename: str) -> list[dict]:
    try:
        content = txt_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        content = txt_path.read_text(encoding="latin-1")
    content = content.strip()
    if not content:
        return []
    entity = clean_entity_evidence(filename.replace("_", " "))
    return [{
        "entity": entity,
        "text": content,
        "chunk_id": f"evidence__{filename}__0000",
        "confidence": None,
    }]


def process_file(file_path: Path, work_dir: Path) -> list[dict]:
    """Process one evidence file, return chunk dicts ready for embedding.

    work_dir: scratch folder for MinerU's markdown output — caller cleans it
    up after use.

    doc/docx/ppt/pptx are converted to PDF first (LibreOffice), then take the PDF path.
    Raises UnsupportedFileError for extensions we don't extract text from
    (jpg/jpeg/png, anything else unrecognized).
    """
    file_path = Path(file_path)
    kind = _file_kind(file_path)
    filename = file_path.stem
    mineru_out_dir = work_dir / "mineru_md"

    if kind == "pdf":
        return _process_pdf(file_path, filename, mineru_out_dir)
    if kind == "convert":
        return _process_pdf(convert_to_pdf(file_path, work_dir / "converted"), filename, mineru_out_dir)
    if kind == "excel":
        return _process_excel(file_path, filename)
    if kind == "csv":
        return _process_csv(file_path, filename)
    if kind == "text":
        return _process_text(file_path, filename)
    raise UnsupportedFileError(f"Tipe file belum didukung untuk ekstraksi: {file_path.name}")
