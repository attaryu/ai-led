from pathlib import Path

import pytest

from ingestion.evidence_pipeline import UnsupportedFileError, _file_kind, process_file


# ── 1 ────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize(
    "name,expected",
    [
        ("a.pdf", "pdf"),
        ("a.PDF", "pdf"),
        ("a.xlsx", "excel"),
        ("a.xls", "excel"),
        ("a.csv", "csv"),
        ("a.txt", "text"),
        ("a.docx", "convert"),
        ("a.PPTX", "convert"),
        ("a.doc", "convert"),
        ("a.jpg", "unknown"),
        ("a.png", "unknown"),
    ],
)
def test_file_kind_by_extension(tmp_path, name, expected):
    p = tmp_path / name
    p.write_bytes(b"not a real file, extension dispatch doesn't read content")
    assert _file_kind(p) == expected


# ── 2 ────────────────────────────────────────────────────────────────────────
def test_file_kind_detects_pdf_via_magic_bytes_without_extension(tmp_path):
    p = tmp_path / "no_extension"
    p.write_bytes(b"%PDF-1.4\n...")
    assert _file_kind(p) == "pdf"


# ── 3 ────────────────────────────────────────────────────────────────────────
def test_process_file_raises_for_unsupported_type(tmp_path):
    p = tmp_path / "photo.jpg"
    p.write_bytes(b"\xff\xd8\xff")  # jpeg magic bytes, still unsupported
    with pytest.raises(UnsupportedFileError):
        process_file(p, tmp_path / "work")


def test_process_file_converts_office_docs_to_pdf_then_takes_the_pdf_path(tmp_path, monkeypatch):
    import ingestion.evidence_pipeline as ep

    src = tmp_path / "SK Rektor.docx"
    src.write_bytes(b"docx bytes")
    converted = tmp_path / "work" / "converted" / "SK Rektor.pdf"
    seen = {}

    def fake_convert(path, out_dir):
        seen["convert"] = (path, out_dir)
        return converted

    def fake_process_pdf(pdf_path, filename, mineru_out_dir):
        seen["pdf"] = (pdf_path, filename)
        return [{"chunk_id": f"evidence__{filename}__0000"}]

    monkeypatch.setattr(ep, "convert_to_pdf", fake_convert)
    monkeypatch.setattr(ep, "_process_pdf", fake_process_pdf)

    chunks = ep.process_file(src, tmp_path / "work")

    assert seen["convert"] == (src, tmp_path / "work" / "converted")
    assert seen["pdf"] == (converted, "SK Rektor")  # chunk ids keep the original file's name
    assert chunks == [{"chunk_id": "evidence__SK Rektor__0000"}]


# ── 4 ────────────────────────────────────────────────────────────────────────
def test_process_file_text(tmp_path):
    p = tmp_path / "catatan.txt"
    p.write_text("Rapat pimpinan menyetujui revisi kurikulum.", encoding="utf-8")
    chunks = process_file(p, tmp_path / "work")
    assert len(chunks) == 1
    assert "revisi kurikulum" in chunks[0]["text"]
    assert chunks[0]["chunk_id"] == "evidence__catatan__0000"


# ── 5 ────────────────────────────────────────────────────────────────────────
def test_process_excel_loads_workbook_once_per_file(tmp_path, monkeypatch):
    """Regression: extract_tables() used to take a filepath and reload+reparse
    the whole workbook per sheet — a 3-sheet file did 3 full loads. Multi-sheet
    LKPS uploads (tens of sheets) made this crawl (see evidence_excel.py)."""
    import openpyxl as openpyxl_module

    from ingestion.evidence_pipeline import _process_excel

    wb = openpyxl_module.Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = 1
    wb.active["B1"] = "x"
    for name in ("Sheet2", "Sheet3"):
        ws = wb.create_sheet(name)
        ws["A1"] = 1
        ws["B1"] = "y"
    p = tmp_path / "multi.xlsx"
    wb.save(p)

    calls = []
    original = openpyxl_module.load_workbook

    def counting_load_workbook(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(openpyxl_module, "load_workbook", counting_load_workbook)

    _process_excel(p, "multi")

    assert len(calls) == 1, (
        f"Expected exactly 1 load_workbook() call for a 3-sheet file, got {len(calls)}"
    )


def test_watermark_and_checkmark_like_dimas():
    import pandas as pd

    from ingestion.evidence_excel import df_to_konten
    from ingestion.watermark import remove_watermark_text

    assert remove_watermark_text("Laporan EDUTECH tahunan") == "Laporan  tahunan"
    assert df_to_konten(pd.DataFrame({"Ada": ["V", "Valid"]})) == "Ada : True\nAda : Valid"


# --- chunking v2: nama dokumen tanpa UUID --------------------------------------------

def test_nama_dokumen_strips_uuid_prefix_and_underscores():
    from ingestion.evidence_pipeline import _nama_dokumen

    uuid_nama = "912b684d-1d94-4c05-90a2-ca983992eac1-13._Pedoman_Kurikulum_2022"
    assert _nama_dokumen(uuid_nama) == "13. Pedoman Kurikulum 2022"
    # tanpa UUID tidak diubah selain garis bawah
    assert _nama_dokumen("SK_Pendirian_PENS") == "SK Pendirian PENS"
    # hanya UUID di awal yang dibuang
    assert _nama_dokumen("laporan-2022-final") == "laporan-2022-final"


def test_process_pdf_keeps_uuid_out_of_chunk_text_but_in_chunk_id(tmp_path, monkeypatch):
    from ingestion import evidence_pipeline as ep

    md = "# BAB I Pendahuluan\n\nIsi pendahuluan yang cukup panjang untuk satu chunk.\n"
    monkeypatch.setattr(ep.mineru_service, "service_enabled", lambda: True)
    monkeypatch.setattr(ep.mineru_service, "parse_pdf", lambda path, **k: md)

    uuid_stem = "912b684d-1d94-4c05-90a2-ca983992eac1-SK_Pendirian"
    hasil = ep._process_pdf(tmp_path / f"{uuid_stem}.pdf", uuid_stem, tmp_path)

    assert hasil, "harus menghasilkan chunk"
    for c in hasil:
        assert "912b684d" not in c["text"]          # tidak ikut ter-embed
        assert "912b684d" not in c["entity"]
        assert set(c) == {"entity", "text", "chunk_id", "confidence"}   # kolom lama tetap
    assert hasil[0]["chunk_id"] == f"evidence__{uuid_stem}__0000"       # unik per berkas
    assert "Dokumen: SK Pendirian" in hasil[0]["text"]


# --- karakter NUL (0x00): Postgres menolaknya ----------------------------------------

def test_bersihkan_nul_removes_every_nul_and_keeps_the_rest():
    from ingestion.evidence_pipeline import _bersihkan_nul

    assert _bersihkan_nul("Formula:\x00\x00\x00 1.5") == "Formula: 1.5"
    assert _bersihkan_nul("tanpa nul") == "tanpa nul"
    assert "\x00" not in _bersihkan_nul("\x00" * 1000)


def test_process_pdf_never_emits_nul_even_when_mineru_output_contains_it(tmp_path, monkeypatch):
    """Kasus nyata 6 Renstra PENS 2020-2024.pdf: 1.160 NUL berderet setelah 'Formula:'."""
    from ingestion import evidence_pipeline as ep

    md = "# BAB I Pendahuluan\n\nIndikator kinerja. Formula:" + "\x00" * 500 + "1\n\nLanjutan teks yang normal.\n"
    monkeypatch.setattr(ep.mineru_service, "service_enabled", lambda: True)
    monkeypatch.setattr(ep.mineru_service, "parse_pdf", lambda path, **k: md)

    hasil = ep._process_pdf(tmp_path / "renstra.pdf", "renstra", tmp_path)

    assert hasil
    assert all("\x00" not in c["text"] and "\x00" not in c["entity"] for c in hasil)
    assert any("Formula:1" in c["text"] for c in hasil)         # teks di sekitar NUL tetap utuh
