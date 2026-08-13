"""
Laporan AK (Kertas Kerja) Excel writer.

Writes a score (col E, index 4) and justification (col H, index 7) to the
JUDGE row identified by its 0-based row_index.

The judge row is the input row (col A = integer indicator number); the Skor
row below it has a formula that copies the value from col E of the judge row.
We write to the judge row's input cells so Excel recalculates on open.
"""

from __future__ import annotations

import threading

import openpyxl

from excel.reader import _COL_JUST, _COL_SCORE, _find_kertas_kerja_sheet

# Serialize concurrent writes to the same file
_write_lock = threading.Lock()


def write_assessment_result(
    xlsx_path: str,
    row_index: int,
    score: int,
    justification: str,
) -> None:
    """Write score and justification to the judge row at row_index (in-place).

    Args:
        xlsx_path:     Path to the Laporan_AK.xlsx file.
        row_index:     0-based row index of the JUDGE row (IndicatorRow.row_index).
        score:         Integer score 1–4.
        justification: Free-text justification string.
    """
    with _write_lock:
        wb = openpyxl.load_workbook(xlsx_path)
        ws = _find_kertas_kerja_sheet(wb)
        if ws is None:
            wb.close()
            raise ValueError(
                f"No Kertas Kerja sheet found in {xlsx_path}; cannot write results."
            )

        # openpyxl rows are 1-based; row_index is 0-based
        openpyxl_row = row_index + 1
        ws.cell(row=openpyxl_row, column=_COL_SCORE + 1).value = score
        ws.cell(row=openpyxl_row, column=_COL_JUST + 1).value = justification

        wb.save(xlsx_path)
        wb.close()
