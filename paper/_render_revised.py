"""Convert PAPER_REVISED_EN.md to DOCX and PDF using python-docx + weasyprint.

No pandoc dependency. Embeds figures from paper/figures/ at the marker locations.

Run from project root:
    python3 paper/_render_revised.py
"""
from __future__ import annotations

import re
from pathlib import Path

import markdown
from weasyprint import CSS, HTML

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
FIGURES = PAPER / "figures"
SRC = PAPER / "PAPER_REVISED_EN.md"
OUT_PDF = PAPER / "PAPER_REVISED_EN.pdf"
OUT_DOCX = PAPER / "PAPER_REVISED_EN.docx"


def load_md() -> str:
    text = SRC.read_text(encoding="utf-8")
    # Strip YAML front-matter
    if text.startswith("---"):
        end = text.find("---", 3)
        if end != -1:
            text = text[end + 3 :].lstrip()
    return text


def replace_figure_markers(md: str) -> str:
    """Replace `**[Figure N here — …]**` markers with markdown image links."""
    figure_map = {
        "Figure 1": "fig1_workflow.png",
        "Figure 2": "fig5_pareto.png",
        "Figure 3": "fig3_per_type_mae.png",
    }
    for label, fname in figure_map.items():
        path = (FIGURES / fname).as_posix()
        # Match: **[Figure N here — caption]**  → markdown image
        pattern = rf"\*\*\[{re.escape(label)} here[^\]]*\]\*\*"
        md = re.sub(pattern, f"\n\n![{label}]({path})\n\n", md)
    return md


def build_html(md_body: str) -> str:
    html_body = markdown.markdown(
        md_body, extensions=["tables", "fenced_code", "attr_list", "md_in_html"]
    )
    css = """
    @page { size: A4; margin: 2.5cm 2cm; }
    body { font-family: 'DejaVu Serif', 'Times New Roman', serif;
           font-size: 11pt; line-height: 1.5; color: #111;
           text-align: justify; }
    h1 { font-size: 16pt; text-align: center; margin: 0 0 12pt; }
    h2 { font-size: 13pt; margin: 18pt 0 8pt; border-bottom: 1pt solid #888; }
    h3 { font-size: 11.5pt; margin: 12pt 0 6pt; font-style: italic; }
    p  { margin: 0 0 8pt; }
    table { border-collapse: collapse; width: 100%; margin: 8pt 0 12pt;
            font-size: 10pt; page-break-inside: avoid; }
    th, td { border: 0.6pt solid #555; padding: 4pt 6pt; vertical-align: top; }
    th { background: #f0f0f0; }
    img { display: block; max-width: 95%; height: auto; margin: 12pt auto;
          page-break-inside: avoid; }
    code { font-family: 'DejaVu Sans Mono', monospace; font-size: 10pt; }
    em, i { font-style: italic; }
    .keywords { font-style: italic; }
    """
    return f"""<!doctype html><html><head><meta charset='utf-8'>
<style>{css}</style></head><body>{html_body}</body></html>"""


def render_pdf(html: str) -> None:
    HTML(string=html, base_url=str(PAPER)).write_pdf(OUT_PDF)
    print(f"  → {OUT_PDF.relative_to(ROOT)}")


def render_docx(md_body: str) -> None:
    """Lightweight DOCX render — paragraphs, headings, tables, images."""
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    # Convert markdown line-by-line (simple state machine).
    lines = md_body.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        # Heading
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level, txt = len(m.group(1)), m.group(2)
            doc.add_heading(_strip_md(txt), level=min(level, 4))
            i += 1
            continue
        # Image
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line)
        if m:
            alt, src = m.group(1), m.group(2)
            try:
                doc.add_picture(src, width=Inches(6.0))
                cap = doc.add_paragraph()
                cap.alignment = 1  # center
                run = cap.add_run(alt)
                run.italic = True
                run.font.size = Pt(9)
            except Exception as e:  # noqa: BLE001
                doc.add_paragraph(f"[Figure missing: {src} ({e})]")
            i += 1
            continue
        # Table (markdown pipe table)
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|$", lines[i + 1]):
            tbl_lines = [line]
            j = i + 2
            while j < len(lines) and lines[j].startswith("|"):
                tbl_lines.append(lines[j])
                j += 1
            _render_table(doc, tbl_lines)
            i = j
            continue
        # Bullet list
        if re.match(r"^[\-\*]\s+", line):
            doc.add_paragraph(_strip_md(re.sub(r"^[\-\*]\s+", "", line)), style="List Bullet")
            i += 1
            continue
        # Numbered list
        if re.match(r"^\d+\.\s+", line):
            doc.add_paragraph(_strip_md(re.sub(r"^\d+\.\s+", "", line)), style="List Number")
            i += 1
            continue
        # Blank
        if not line.strip():
            i += 1
            continue
        # Paragraph (may span multiple lines)
        para_lines = [line]
        j = i + 1
        while j < len(lines) and lines[j].strip() and not _is_block_start(lines[j]):
            para_lines.append(lines[j])
            j += 1
        doc.add_paragraph(_strip_md(" ".join(para_lines)))
        i = j

    doc.save(OUT_DOCX)
    print(f"  → {OUT_DOCX.relative_to(ROOT)}")


def _is_block_start(line: str) -> bool:
    return bool(
        re.match(r"^(#{1,4}\s|\!\[|\||[\-\*]\s|\d+\.\s)", line.lstrip())
    )


def _strip_md(s: str) -> str:
    s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)  # bold
    s = re.sub(r"\*([^*]+)\*", r"\1", s)  # italic
    s = re.sub(r"`([^`]+)`", r"\1", s)  # code
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)  # link
    return s


def _render_table(doc, tbl_lines: list[str]) -> None:
    from docx.shared import Pt

    rows = []
    for line in tbl_lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(n_cols):
            cell_text = _strip_md(row[ci] if ci < len(row) else "")
            cell = table.rows[ri].cells[ci]
            cell.text = ""
            run = cell.paragraphs[0].add_run(cell_text)
            run.font.size = Pt(10)
            if ri == 0:
                run.bold = True


def main() -> None:
    print(f"Source: {SRC.relative_to(ROOT)}")
    md = load_md()
    md = replace_figure_markers(md)
    print("Rendering DOCX…")
    render_docx(md)
    print("Rendering PDF…")
    html = build_html(md)
    render_pdf(html)
    print("\n✓ Done.")


if __name__ == "__main__":
    main()
